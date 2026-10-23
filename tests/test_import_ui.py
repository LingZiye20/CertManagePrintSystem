"""Excel 导入界面测试：文件选择、结果展示与错误明细。"""

from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from PySide6.QtWidgets import QApplication, QFileDialog, QMenu, QPushButton

from app.bootstrap import bootstrap
from app.config.paths import AppPaths
from app.services import ExcelImportService, ImportIssue, ImportResult, StudentService
from app.ui.import_result_dialog import ImportResultDialog
from app.ui.main_window import MainWindow
from app.ui.student_page import StudentPage
from app.utils.excel_template import EXCEL_COLUMN_NAMES


@pytest.fixture(scope="session")
def qapp() -> QApplication:
    return QApplication.instance() or QApplication([])


def _write_students(path: Path, student_numbers: list[str | None]) -> Path:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "学生数据"
    sheet.append(list(EXCEL_COLUMN_NAMES))
    for student_no in student_numbers:
        row = [None] * 37
        row[13] = student_no
        row[14] = "测试"
        sheet.append(row)
    workbook.save(path)
    workbook.close()
    return path


@pytest.fixture
def page(
    student_service: StudentService, qapp: QApplication
) -> StudentPage:
    import_service = ExcelImportService(student_service.database)
    widget = StudentPage(student_service, import_service=import_service)
    widget.show()
    qapp.processEvents()
    yield widget
    widget.close()


def test_import_menu_action_is_available(
    tmp_path: Path, qapp: QApplication
) -> None:
    from PySide6.QtGui import QAction

    context = bootstrap(AppPaths.for_program(tmp_path))
    window = MainWindow(context=context)
    window.show()
    qapp.processEvents()

    page = window.student_page
    assert page is not None
    assert page.import_service is not None
    actions = {action.text(): action for action in window.findChildren(QAction)}
    assert "导入 Excel(&I)" in actions
    assert actions["导入 Excel(&I)"].isEnabled() is True
    assert "下载导入模板(&T)…" in actions
    assert actions["下载导入模板(&T)…"].isEnabled() is True
    assert "新增记录(&A)" not in actions
    assert not hasattr(page, "add_button")
    assert not hasattr(page, "import_button")
    window.close()


def test_data_menu_order_and_toolbar_without_add_entry(
    tmp_path: Path, qapp: QApplication
) -> None:
    context = bootstrap(AppPaths.for_program(tmp_path))
    window = MainWindow(context=context)
    window.show()
    qapp.processEvents()

    data_menu = next(
        menu for menu in window.findChildren(QMenu) if menu.title() == "数据管理(&D)"
    )
    assert [action.text() for action in data_menu.actions()] == [
        "下载导入模板(&T)…",
        "导入 Excel(&I)",
        "导出数据(&E)…",
        "",
        "刷新(&R)",
    ]

    # 编辑/删除已从数据管理菜单移除
    action_texts = {a.text() for a in data_menu.actions()}
    assert "编辑记录(&E)" not in action_texts
    assert "删除记录(&D)" not in action_texts

    page = window.student_page
    assert page is not None
    toolbar_texts = [button.text() for button in page.findChildren(QPushButton)]
    assert "新增" not in toolbar_texts
    assert "导入 Excel" not in toolbar_texts
    window.close()


def test_cancel_file_dialog_does_not_write_or_show_success(
    page: StudentPage, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(QFileDialog, "getOpenFileName", lambda *args, **kwargs: ("", ""))
    called = {"value": False}

    def _should_not_run(path):
        called["value"] = True
        raise AssertionError("取消选择后不应调用导入服务")

    monkeypatch.setattr(page.import_service, "import_file", _should_not_run)

    page.import_excel()

    assert called["value"] is False
    assert page.service.records.count_records() == 0
    assert "成功" not in page.message_label.text()


def test_successful_import_updates_message_and_table(
    page: StudentPage,
    tmp_path: Path,
    qapp: QApplication,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    path = _write_students(tmp_path / "students.xlsx", ["00123", "00124"])
    monkeypatch.setattr(
        QFileDialog, "getOpenFileName", lambda *args, **kwargs: (str(path), "")
    )
    monkeypatch.setattr(ImportResultDialog, "exec", lambda self: 0)

    page.import_excel()
    qapp.processEvents()

    assert "成功" in page.message_label.text()
    assert page.service.records.count_records() == 2
    assert page.model.rowCount() == 2


def test_failed_import_does_not_show_success_or_insert_rows(
    page: StudentPage,
    tmp_path: Path,
    qapp: QApplication,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    path = _write_students(tmp_path / "students.xlsx", [None])
    monkeypatch.setattr(
        QFileDialog, "getOpenFileName", lambda *args, **kwargs: (str(path), "")
    )
    monkeypatch.setattr(ImportResultDialog, "exec", lambda self: 0)

    page.import_excel()
    qapp.processEvents()

    assert "成功" not in page.message_label.text()
    assert "导入失败" in page.message_label.text()
    assert page.service.records.count_records() == 0


def test_success_is_not_shown_while_import_is_running(
    page: StudentPage,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    path = _write_students(tmp_path / "students.xlsx", ["00123"])
    monkeypatch.setattr(
        QFileDialog, "getOpenFileName", lambda *args, **kwargs: (str(path), "")
    )
    monkeypatch.setattr(ImportResultDialog, "exec", lambda self: 0)
    original = page.import_service.import_file

    def _observed(path_arg):
        assert "成功" not in page.message_label.text()
        return original(path_arg)

    monkeypatch.setattr(page.import_service, "import_file", _observed)

    page.import_excel()

    assert "成功" in page.message_label.text()


def test_result_dialog_lists_row_column_field_and_reason(
    qapp: QApplication,
) -> None:
    result = ImportResult(
        file_name="students.xlsx",
        status="failed",
        total_rows=1,
        imported_count=0,
        failed_count=1,
        issues=(
            ImportIssue(
                row=2,
                column="学号",
                field="student_no",
                kind="required",
                reason="学号不能为空",
            ),
        ),
        batch_id=7,
        batch_saved=True,
        warnings=("失败审计记录未能保存，请查看日志",),
        message="导入失败：未写入任何学生记录（1 行未导入）",
    )

    dialog = ImportResultDialog(result)

    assert "总行数：1" in dialog.summary_label.text()
    assert "失败：1" in dialog.summary_label.text()
    assert dialog.issues_table is not None
    assert dialog.issues_table.rowCount() == 1
    assert dialog.issues_table.item(0, 0).text() == "2"
    assert dialog.issues_table.item(0, 1).text() == "学号"
    assert dialog.issues_table.item(0, 2).text() == "student_no"
    assert dialog.issues_table.item(0, 3).text() == "学号不能为空"
    assert "失败审计记录未能保存" in dialog.warnings_label.text()
    dialog.close()


def test_result_dialog_success_has_no_error_table(qapp: QApplication) -> None:
    result = ImportResult(
        file_name="students.xlsx",
        status="success",
        total_rows=2,
        imported_count=2,
        failed_count=0,
        issues=(),
        batch_id=8,
        batch_saved=True,
        warnings=(),
        message="导入成功：2 条",
    )

    dialog = ImportResultDialog(result)

    assert "总行数：2" in dialog.summary_label.text()
    assert "成功：2" in dialog.summary_label.text()
    assert dialog.issues_table is None
    dialog.close()


def test_main_window_injects_import_service(
    tmp_path: Path, qapp: QApplication
) -> None:
    context = bootstrap(AppPaths.for_program(tmp_path))
    window = MainWindow(context=context)

    assert window.student_page is not None
    assert window.student_page.import_service is not None
    window.close()


def test_download_template_writes_header_only_file(
    page: StudentPage,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    from app.ui import student_page as student_page_module

    target = tmp_path / "证书导入模板.xlsx"
    messages: list[str] = []
    monkeypatch.setattr(
        QFileDialog,
        "getSaveFileName",
        staticmethod(lambda *args, **kwargs: (str(target), "")),
    )
    monkeypatch.setattr(
        student_page_module.QMessageBox,
        "information",
        staticmethod(
            lambda parent, title, message, *args, **kwargs: messages.append(message)
        ),
    )

    page.download_import_template()

    assert target.is_file()
    assert messages and str(target) in messages[0]
    assert str(target) in page.message_label.text()

    workbook = load_workbook(target, read_only=True)
    try:
        rows = list(workbook.active.iter_rows(values_only=True))
    finally:
        workbook.close()
    assert len(rows) == 1
    assert tuple(rows[0]) == EXCEL_COLUMN_NAMES


def test_download_template_cancel_writes_nothing(
    page: StudentPage,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(
        QFileDialog,
        "getSaveFileName",
        staticmethod(lambda *args, **kwargs: ("", "")),
    )

    page.download_import_template()

    assert list(tmp_path.glob("*.xlsx")) == []
    assert page.message_label.text() == ""


def test_download_template_failure_removes_partial_file(
    page: StudentPage,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    from app.ui import student_page as student_page_module

    target = tmp_path / "broken.xlsx"
    warnings: list[str] = []

    def _broken(path):
        Path(path).write_bytes(b"partial")
        raise OSError("disk full")

    monkeypatch.setattr(student_page_module, "write_template", _broken)
    monkeypatch.setattr(
        QFileDialog,
        "getSaveFileName",
        staticmethod(lambda *args, **kwargs: (str(target), "")),
    )
    monkeypatch.setattr(
        student_page_module.QMessageBox,
        "warning",
        staticmethod(
            lambda parent, title, message, *args, **kwargs: warnings.append(message)
        ),
    )

    page.download_import_template()

    assert not target.exists()
    assert warnings and "无法写入" in warnings[0]
