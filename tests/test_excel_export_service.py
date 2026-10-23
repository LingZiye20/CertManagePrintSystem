"""固定 37 列 Excel 数据导出 Service 测试。"""

from __future__ import annotations

from pathlib import Path
from unittest.mock import patch

import openpyxl
import pytest

from app.data.db import EXCEL_FIELD_NAMES
from app.services import (
    AuditWriteError,
    ExcelExportService,
    ExcelImportService,
    ExportResult,
    StudentQuery,
)
from app.utils.excel_template import EXCEL_COLUMN_NAMES, EXCEL_SHEET_TITLE, write_template

CERTIFICATE_FIELDS = (
    "certificate_no",
    "degree_certificate_no",
    "other_certificate_no",
)


@pytest.fixture
def export_service(student_service, repository_database) -> ExcelExportService:
    return ExcelExportService(repository_database, student_service=student_service)


def _create(student_service, student_no: str = "001", **overrides) -> int:
    data = {name: None for name in EXCEL_FIELD_NAMES}
    data["student_no"] = student_no
    data["name"] = f"学生{student_no}"
    data.update(overrides)
    return student_service.create_record(data)


def _load_rows(path: Path) -> list[tuple]:
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.worksheets[0]
        return list(sheet.iter_rows(values_only=True))
    finally:
        workbook.close()


# ---------------------------------------------------------------- 基础结构


def test_export_header_matches_excel_column_names(
    export_service: ExcelExportService, student_service, tmp_path: Path
) -> None:
    _create(student_service, student_no="001")
    target = tmp_path / "out.xlsx"

    result = export_service.export(StudentQuery(), target)

    assert isinstance(result, ExportResult)
    rows = _load_rows(target)
    assert len(rows) >= 1
    assert tuple(rows[0]) == EXCEL_COLUMN_NAMES
    assert len(rows[0]) == 37


def test_export_sheet_title_is_student_data(
    export_service: ExcelExportService, student_service, tmp_path: Path
) -> None:
    _create(student_service)
    target = tmp_path / "out.xlsx"

    export_service.export(StudentQuery(), target)

    workbook = openpyxl.load_workbook(target, read_only=True)
    try:
        assert workbook.worksheets[0].title == EXCEL_SHEET_TITLE
    finally:
        workbook.close()


def test_export_row_count_matches_all_records(
    export_service: ExcelExportService, student_service, tmp_path: Path
) -> None:
    _create(student_service, "001")
    _create(student_service, "002")
    _create(student_service, "003")
    target = tmp_path / "out.xlsx"

    result = export_service.export(StudentQuery(), target)

    assert result.record_count == 3
    rows = _load_rows(target)
    assert len(rows) == 1 + 3


# ---------------------------------------------------------------- 文本保真


def test_leading_zero_student_no_preserved_as_text(
    export_service: ExcelExportService, student_service, tmp_path: Path
) -> None:
    _create(student_service, student_no="00123")
    target = tmp_path / "out.xlsx"

    export_service.export(StudentQuery(), target)

    workbook = openpyxl.load_workbook(target)
    try:
        sheet = workbook.active
        # 学号在 EXCEL_FIELD_NAMES 中的位置（索引从 1 开始）
        col_index = EXCEL_FIELD_NAMES.index("student_no") + 1
        value = sheet.cell(row=2, column=col_index).value
    finally:
        workbook.close()

    assert value == "00123"
    assert isinstance(value, str)


def test_certificate_numbers_preserved_verbatim(
    export_service: ExcelExportService, student_service, tmp_path: Path
) -> None:
    _create(
        student_service,
        student_no="001",
        certificate_no="CERT-001",
        degree_certificate_no="DEG-001",
        other_certificate_no="OTHER-001",
    )
    target = tmp_path / "out.xlsx"

    export_service.export(StudentQuery(), target)

    workbook = openpyxl.load_workbook(target)
    try:
        sheet = workbook.active
        row_values = {
            name: sheet.cell(row=2, column=idx + 1).value
            for idx, name in enumerate(EXCEL_FIELD_NAMES)
        }
    finally:
        workbook.close()

    assert row_values["certificate_no"] == "CERT-001"
    assert row_values["degree_certificate_no"] == "DEG-001"
    assert row_values["other_certificate_no"] == "OTHER-001"


def test_none_fields_written_as_empty_string(
    export_service: ExcelExportService, student_service, tmp_path: Path
) -> None:
    _create(student_service, student_no="001")  # 其余字段均为 None
    target = tmp_path / "out.xlsx"

    export_service.export(StudentQuery(), target)

    workbook = openpyxl.load_workbook(target)
    try:
        sheet = workbook.active
        for idx in range(len(EXCEL_FIELD_NAMES)):
            if EXCEL_FIELD_NAMES[idx] in ("student_no", "name"):
                continue
            cell = sheet.cell(row=2, column=idx + 1)
            # openpyxl 把空字符串读回为 None，二者都视为空值
            assert cell.value in (None, "")
    finally:
        workbook.close()


# ---------------------------------------------------------------- 筛选与空结果


def test_empty_result_exports_header_only(
    export_service: ExcelExportService, student_service, tmp_path: Path
) -> None:
    _create(student_service, "001")
    target = tmp_path / "out.xlsx"

    result = export_service.export(
        StudentQuery(filters={"student_no": "不存在的学号"}), target
    )

    assert result.record_count == 0
    rows = _load_rows(target)
    assert len(rows) == 1
    assert tuple(rows[0]) == EXCEL_COLUMN_NAMES


def test_filter_exports_subset(
    export_service: ExcelExportService, student_service, tmp_path: Path
) -> None:
    _create(student_service, "001", name="张三")
    _create(student_service, "002", name="李四")
    _create(student_service, "003", name="张三")
    target = tmp_path / "out.xlsx"

    result = export_service.export(StudentQuery(filters={"name": "张三"}), target)

    assert result.record_count == 2
    workbook = openpyxl.load_workbook(target)
    try:
        sheet = workbook.active
        names = [
            sheet.cell(row=r, column=EXCEL_FIELD_NAMES.index("name") + 1).value
            for r in range(2, sheet.max_row + 1)
        ]
    finally:
        workbook.close()
    assert names == ["张三", "张三"]


# ---------------------------------------------------------------- 审计


def test_success_audit_written_without_full_id_number(
    export_service: ExcelExportService, student_service, tmp_path: Path
) -> None:
    secret = "ID-SECRET-123456"
    _create(student_service, "001", id_number=secret)
    target = tmp_path / "out.xlsx"

    export_service.export(StudentQuery(filters={"name": "学生001"}), target)

    logs = export_service.logs.list_logs(operation_type="excel_export")
    assert len(logs) == 1
    log = logs[0]
    assert log["result"] == "success"
    assert secret not in (log["summary"] or "")
    assert "学生001" not in (log["summary"] or "")  # 筛选值不写入摘要
    assert "1 条" in (log["summary"] or "")


def test_success_audit_summary_contains_file_name_and_count(
    export_service: ExcelExportService, student_service, tmp_path: Path
) -> None:
    _create(student_service, "001")
    target = tmp_path / "custom-name.xlsx"

    export_service.export(StudentQuery(), target)

    log = export_service.logs.list_logs(operation_type="excel_export")[0]
    assert "custom-name.xlsx" in (log["summary"] or "")
    assert "1 条" in (log["summary"] or "")


def test_success_audit_failure_raises_audit_write_error(
    export_service: ExcelExportService, student_service, tmp_path: Path
) -> None:
    _create(student_service, "001")
    target = tmp_path / "out.xlsx"

    with patch.object(
        export_service.logs, "create", side_effect=__import__("sqlite3").OperationalError("down")
    ):
        with pytest.raises(AuditWriteError):
            export_service.export(StudentQuery(), target)

    # 文件仍然生成（导出是真实产物）
    assert target.exists()
    rows = _load_rows(target)
    assert len(rows) == 2


def test_oserror_writes_failed_audit_and_removes_partial_file(
    export_service: ExcelExportService, student_service, tmp_path: Path
) -> None:
    _create(student_service, "001")
    target = tmp_path / "out.xlsx"
    assert not target.exists()

    with patch(
        "app.services.excel_export_service.Workbook.save",
        side_effect=OSError("disk full"),
    ):
        with pytest.raises(Exception) as exc_info:
            export_service.export(StudentQuery(), target)

    assert "导出" in str(exc_info.value)
    # 半截文件被清理
    assert not target.exists()
    # 失败审计写入
    logs = export_service.logs.list_logs(operation_type="excel_export")
    assert len(logs) == 1
    assert logs[0]["result"] == "failure"
    assert "out.xlsx" in (logs[0]["summary"] or "")


# ---------------------------------------------------------------- 闭环兼容


def test_export_columns_match_template_structure(
    export_service: ExcelExportService, student_service, tmp_path: Path
) -> None:
    _create(student_service, "001", certificate_no="C-1")
    exported = tmp_path / "exported.xlsx"
    template = tmp_path / "template.xlsx"

    export_service.export(StudentQuery(), exported)
    write_template(template)

    export_rows = _load_rows(exported)
    template_rows = _load_rows(template)

    assert tuple(export_rows[0]) == tuple(template_rows[0])
    assert len(export_rows[0]) == len(template_rows[0]) == 37


def test_exported_file_passes_import_structure_validation(
    export_service: ExcelExportService, student_service, repository_database, tmp_path: Path
) -> None:
    _create(student_service, "001", certificate_no="C-1")
    exported = tmp_path / "exported.xlsx"

    export_service.export(StudentQuery(), exported)

    # 清空三类证书编号列，避免导入时因证书编号非空被拒
    workbook = openpyxl.load_workbook(exported)
    try:
        sheet = workbook.active
        for field in CERTIFICATE_FIELDS:
            col = EXCEL_FIELD_NAMES.index(field) + 1
            for row in range(2, sheet.max_row + 1):
                sheet.cell(row=row, column=col).value = ""
        workbook.save(exported)
    finally:
        workbook.close()

    result = ExcelImportService(repository_database).import_file(exported)

    reasons = " ".join(issue.reason for issue in result.issues)
    kinds = {issue.kind for issue in result.issues}
    assert "列名" not in reasons
    assert "列数" not in reasons
    assert "列顺序" not in reasons
    assert "template" not in kinds
