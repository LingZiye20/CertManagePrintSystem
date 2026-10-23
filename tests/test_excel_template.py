"""Excel 导入模板生成测试：表头、格式与导入校验兼容性。"""

from __future__ import annotations

from pathlib import Path

import openpyxl

from app.services import ExcelImportService
from app.utils.excel_template import (
    EXCEL_COLUMN_NAMES,
    EXCEL_SHEET_TITLE,
    write_template,
)


def test_write_template_contains_only_header_row(tmp_path: Path) -> None:
    path = tmp_path / "template.xlsx"

    write_template(path)

    workbook = openpyxl.load_workbook(path, read_only=True)
    try:
        assert len(workbook.worksheets) == 1
        sheet = workbook.worksheets[0]
        assert sheet.title == EXCEL_SHEET_TITLE
        rows = list(sheet.iter_rows(values_only=True))
    finally:
        workbook.close()

    assert len(rows) == 1
    assert tuple(rows[0]) == EXCEL_COLUMN_NAMES
    assert len(rows[0]) == 37


def test_template_header_is_bold_and_frozen(tmp_path: Path) -> None:
    path = tmp_path / "template.xlsx"

    write_template(path)

    workbook = openpyxl.load_workbook(path)
    try:
        sheet = workbook.active
        assert sheet.freeze_panes == "A2"
        assert all(cell.font.bold for cell in sheet[1])
    finally:
        workbook.close()


def test_template_passes_import_structure_validation(
    repository_database, tmp_path: Path
) -> None:
    path = tmp_path / "template.xlsx"
    write_template(path)

    result = ExcelImportService(repository_database).import_file(path)

    assert all(issue.kind != "template" for issue in result.issues)
    reasons = " ".join(issue.reason for issue in result.issues)
    assert "列名" not in reasons
    assert "列数" not in reasons
    assert "列顺序" not in reasons
