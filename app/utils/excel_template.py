"""固定 Excel 模板列名与模板文件生成。

列名来自提示词包 C.2 的 Excel 映射表，顺序与
:data:`app.data.db.EXCEL_FIELD_NAMES` 严格一一对应。
界面标签（``StudentDialog.FIELD_LABELS``）不是 Excel 表头，不得复用。
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from app.data.db import EXCEL_FIELD_NAMES

EXCEL_HEADER_ROW = 1
EXCEL_COLUMN_COUNT = 37
EXCEL_SHEET_TITLE = "学生数据"

EXCEL_COLUMN_NAMES: tuple[str, ...] = (
    "毕结业批次",
    "毕业审核",
    "学位审核",
    "学生类型",
    "院号",
    "学位类别码",
    "学位类别全",
    "专业类别",
    "打印标记",
    "证书编号",
    "学位证书编号",
    "考生号",
    "省份代码",
    "学号",
    "姓名",
    "性别",
    "出生日期",
    "证件类型",
    "证件号码",
    "政治面貌",
    "民族",
    "院校代码",
    "院校名称",
    "专业代码",
    "专业名称",
    "分院",
    "学院名",
    "班级",
    "层次",
    "学制",
    "学习形式",
    "当前年级",
    "入学日期",
    "毕业日期",
    "注册状态",
    "备注",
    "其他证书编号",
)

if len(EXCEL_COLUMN_NAMES) != len(EXCEL_FIELD_NAMES):  # pragma: no cover - 开发期保护
    raise RuntimeError("Excel 模板列数与 EXCEL_FIELD_NAMES 不一致")


def write_template(path: str | Path) -> None:
    """生成仅含表头行的固定 37 列导入模板。

    表头逐字使用 :data:`EXCEL_COLUMN_NAMES`，不写入任何数据行或示例行。
    表头加粗、冻结首行并按列名粗略设置列宽，均不影响导入兼容性。
    目录不可写或磁盘错误等 ``OSError`` 原样抛出，由界面层转换提示。
    """
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = EXCEL_SHEET_TITLE
    sheet.append(list(EXCEL_COLUMN_NAMES))

    header_row = sheet[EXCEL_HEADER_ROW]
    for cell in header_row:
        cell.font = Font(bold=True)
    sheet.freeze_panes = "A2"
    for index, cell in enumerate(header_row, start=1):
        label = str(cell.value or "")
        sheet.column_dimensions[get_column_letter(index)].width = max(
            10, min(24, len(label) * 2 + 4)
        )

    workbook.save(str(path))
